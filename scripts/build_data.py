#!/usr/bin/env python3
"""Ops dashboard data builder.

Downloads the three public Google Drive folders (no API key), recognises every
Excel file by its CONTENT (never by filename) and writes data/data.json.

Usage:
  python scripts/build_data.py                 # download from Drive, then build
  python scripts/build_data.py --local DIR     # use DIR/tilldate, DIR/monthend, DIR/performance

Exit code 1 = critical problem; data.json is NOT overwritten, so the live
dashboard keeps showing the last good data.
"""
import argparse, calendar, datetime as dt, glob, json, os, re, shutil, sys, tempfile
from openpyxl import load_workbook

FOLDERS = {
    "tilldate": "12UEEFUoUIP_duIJGYMw5g2zjsiuphjOa",
    "monthend": "1oRRVEZ7A6AjrsXZ-wHx7JQrDh3FoB2eb",
    "performance": "1r09IGv8Pk0J86li4hqLD8gskj5j3SBih",
}
OUT = os.environ.get("DATA_OUT") or os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "data", "data.json")
CODE_RE = re.compile(r"^[A-Z]{1,2}\d{2,4}$")
MONTH_RE = re.compile(r"^([A-Za-z]{3})'(\d{2})$")
MONTHS = {m.lower(): i for i, m in enumerate(calendar.month_abbr) if m}
ISSUES = []


def issue(level, source, msg):
    ISSUES.append({"level": level, "source": source, "message": msg})
    print(f"[{level}] {source}: {msg}")


def norm(v):
    return re.sub(r"\s+", " ", str(v)).strip().lower() if v is not None else ""


def num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v) if v == v and abs(v) != float("inf") else None
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return None


def r(v, d=2):
    return None if v is None else round(v, d)


def as_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)) and 30000 < v < 60000:  # Excel serial
        return dt.date(1899, 12, 30) + dt.timedelta(days=int(v))
    return None


def header_map(row):
    """First occurrence of each normalised header -> column index."""
    m = {}
    for i, v in enumerate(row):
        k = norm(v)
        if k and k not in m:
            m[k] = i
    return m


def pick(m, *names):
    for n in names:
        if n in m:
            return m[n]
    for n in names:  # prefix fallback, e.g. "outlet incharge name"
        for k, i in m.items():
            if k.startswith(n):
                return i
    return None


# --------------------------------------------------------------------------- download
def download(workdir):
    import gdown
    for key, fid in FOLDERS.items():
        out = os.path.join(workdir, key)
        os.makedirs(out, exist_ok=True)
        try:
            res = gdown.download_folder(id=fid, output=out, quiet=True, use_cookies=False, remaining_ok=True)
            if not res:
                issue("error", key, "Drive folder returned no files. Check it is shared as 'Anyone with the link' and holds at most 50 files.")
        except Exception as e:  # keep going; missing folders become issues
            issue("error", key, f"Could not download Drive folder ({e}). Check it is shared as 'Anyone with the link'.")


def excel_files(folder):
    files = []
    for p in glob.glob(os.path.join(folder, "**", "*"), recursive=True):
        b = os.path.basename(p)
        if os.path.isfile(p) and b.lower().endswith((".xlsx", ".xlsm")) and not b.startswith("~$"):
            files.append(p)
    return sorted(files)


# --------------------------------------------------------------------------- recognition
def first_rows(ws, n=4, cols=120):
    return [list(r) for r in ws.iter_rows(min_row=1, max_row=n, max_col=cols, values_only=True)]


def classify(path):
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        issue("error", os.path.basename(path), f"Not a readable Excel file ({e}).")
        return None, None
    heads = {}
    for ws in wb.worksheets:
        try:
            heads[ws.title] = first_rows(ws)
        except Exception:
            continue
    flat = lambda rows: {norm(v) for row in rows for v in row if v is not None}
    for rows in heads.values():
        f = flat(rows)
        if "head name" in f and "calc type" in f:
            return "kpi", (wb, heads)
    for rows in heads.values():
        f = flat(rows)
        if {"sales this", "ff this", "bs this"} <= f:
            return "business", (wb, heads)
    for rows in heads.values():
        f = flat(rows)
        if {"code", "sales revenue", "total outlet level opex"} <= f:
            return "pnl", (wb, heads)
    for rows in heads.values():
        f = flat(rows)
        if "code" in f and "zonal" in f and "format" in f:
            return "master", (wb, heads)
    issue("warn", os.path.basename(path), "Unrecognised file: no known sheet layout found. It was skipped.")
    return None, None


# --------------------------------------------------------------------------- business report
def version(title):
    m = re.search(r"\(v(\d+)\)", title, re.I)
    return int(m.group(1)) if m else 0


def find_header_row(rows, must):
    for i, row in enumerate(rows):
        f = {norm(v) for v in row if v is not None}
        if must <= f:
            return i
    return None


def parse_outlet_sheet(ws, hdr_idx):
    rows = ws.iter_rows(values_only=True)
    out, hm = {}, None
    for i, row in enumerate(rows):
        if i < hdr_idx:
            continue
        if i == hdr_idx:
            hm = header_map(row)
            c = {
                "code": pick(hm, "code"), "name": pick(hm, "all store", "outlet name"),
                "rl": pick(hm, "leader"), "zn": pick(hm, "zonal head", "zonal"),
                "oi": pick(hm, "outlet incharge"),
                "s": pick(hm, "sales this"), "s0": pick(hm, "sales last"),
                "f": pick(hm, "ff this"), "f0": pick(hm, "ff last"),
                "b": pick(hm, "bs this"), "b0": pick(hm, "bs last"),
                "gp": pick(hm, "gp% this"), "gp0": pick(hm, "gp% last"),
                "gv": pick(hm, "gpv this"), "gv0": pick(hm, "gpv last"),
                "cv": pick(hm, "cpv this"), "cv0": pick(hm, "cpv last"),
            }
            continue
        code = str(row[c["code"]]).strip() if c["code"] is not None and row[c["code"]] is not None else ""
        if not CODE_RE.match(code) or code in out:
            continue
        g = lambda k: row[c[k]] if c[k] is not None and c[k] < len(row) else None
        rec = {"name": g("name"), "rl": g("rl"), "zn": g("zn"), "oi": g("oi")}
        for k in ("s", "s0", "f", "f0", "b", "b0", "gv", "gv0", "cv", "cv0"):
            rec[k] = num(g(k))
        for k in ("gp", "gp0"):
            rec[k] = num(g(k))
        out[code] = rec
    return out


def parse_business(path, wb, heads):
    src = os.path.basename(path)
    rep = {"file": src}
    # Report date + targets: sheet whose A1 is a date and row 1 mentions "target for operations"
    tgt_sheet = None
    for t, rows in heads.items():
        if rows and as_date(rows[0][0] if rows[0] else None) and any("target for operations" in norm(v) for v in rows[0]):
            tgt_sheet = t
            break
    if not tgt_sheet:
        issue("error", src, "Target sheet (date in A1, 'Target for Operations') not found.")
        return None
    rep["date"] = as_date(heads[tgt_sheet][0][0])
    targets = {}
    hi = find_header_row(heads[tgt_sheet], {"code", "targ", "achv"})
    if hi is None:
        issue("error", src, f"Sheet '{tgt_sheet}': headers CODE / TARG / ACHV not found.")
    else:
        hm = header_map(heads[tgt_sheet][hi])
        ci, ti, ai = pick(hm, "code"), pick(hm, "targ"), pick(hm, "achv")
        ni, li, zi = pick(hm, "outlet name"), pick(hm, "leader"), pick(hm, "zonal")
        for i, row in enumerate(wb[tgt_sheet].iter_rows(values_only=True)):
            if i <= hi or ci >= len(row):
                continue
            code = str(row[ci]).strip() if row[ci] else ""
            if CODE_RE.match(code) and code not in targets:
                cell = lambda j: str(row[j]).strip() if j is not None and j < len(row) and row[j] else None
                targets[code] = (num(row[ti]), num(row[ai]), cell(ni), cell(li), cell(zi))

    # Outlet comparison sheets (SPLY / SPLM) recognised by headers + title text
    cands = {"SPLY": [], "SPLM": []}
    for t, rows in heads.items():
        hi = find_header_row(rows, {"code", "sales this", "sales last", "ff this", "bs this"})
        if hi is None:
            continue
        title = " ".join(str(v) for row in rows[:hi] for v in row if isinstance(v, str))
        kind = "SPLY" if "SPLY" in title.upper() else "SPLM" if "SPLM" in title.upper() else None
        if not kind:
            continue
        label = next((str(v).strip() for row in rows[:hi] for v in row if isinstance(v, str) and "same day" in v.lower()), "")
        cands[kind].append((t, hi, re.sub(r"\s+", " ", label)))

    def choose(kind, same):
        pool = []
        for t, hi, label in cands[kind]:
            up = t.upper()
            if "(2)" in up or "OWN" in up or "FRAN" in up:
                continue
            if same != ("SAME" in up):
                continue
            pool.append((version(t), t, hi, label))
        return max(pool) if pool else None

    outlets = {}
    for kind in ("SPLY", "SPLM"):
        ch = choose(kind, False)
        if not ch:
            issue("error", src, f"No all-store {kind} outlet sheet found.")
            continue
        _, t, hi, label = ch
        rep[kind.lower() + "_label"] = label
        rep[kind.lower() + "_sheet"] = t
        data = parse_outlet_sheet(wb[t], hi)
        suffix = "y" if kind == "SPLY" else "m"
        for code, d in data.items():
            o = outlets.setdefault(code, {"c": code})
            for k in ("name", "rl", "zn", "oi"):
                if d[k] and not o.get(k):
                    o[k] = str(d[k]).strip()
            for k in ("s", "f", "b", "gp", "gv", "cv"):
                if o.get(k) is None and d[k] is not None:
                    o[k] = d[k]
                o[k + suffix] = d[k + "0"]
        ss = choose(kind, True)
        if ss:
            for code in parse_outlet_sheet(wb[ss[1]], ss[2]):
                if code in outlets:
                    outlets[code]["ss" + suffix] = 1
            rep["same_" + kind.lower() + "_sheet"] = ss[1]
        else:
            issue("warn", src, f"No same-store {kind} sheet found; same-store view unavailable for {kind}.")

    for code, (t, a, nm, rl, zn) in targets.items():
        o = outlets.setdefault(code, {"c": code})
        o["t"], o["a"] = t, a
        for k, v in (("name", nm), ("rl", rl), ("zn", zn)):
            if v and not o.get(k):
                o[k] = v
        if o.get("s") is None:
            o["s"] = a
    rep["categories"] = parse_categories(wb, heads, src)

    # compact + round
    clean = []
    for o in outlets.values():
        for k in list(o):
            if isinstance(o[k], float):
                o[k] = r(o[k], 4 if k.startswith("gp") else 0 if k[0] in "sfgcta" else 2)
            if o[k] is None or o[k] == "":
                del o[k]
        clean.append(o)
    rep["outlets"] = clean
    if len(clean) < 100:
        issue("error", src, f"Only {len(clean)} outlets found; expected about 1,000.")
    return rep


def parse_categories(wb, heads, src):
    """Category growth blocks ('Same Day SPLM/SPLY ...' then a header row with 'Sales Growth')."""
    for t in heads:
        top = [list(x) for x in wb[t].iter_rows(max_row=30, max_col=2, values_only=True)]
        if not any("same day" in norm(row[0]) for row in top if row):
            continue
        allrows = [list(x) for x in wb[t].iter_rows(max_row=80, max_col=30, values_only=True)]
        blocks = {}
        i = 0
        while i < len(allrows) - 1:
            title = norm(allrows[i][0]) if allrows[i] else ""
            if title.startswith("same day") and any("sales growth" in norm(v) for v in allrows[i + 1]):
                kind = "sply" if "sply" in title else "splm"
                grp = {}
                for j, v in enumerate(allrows[i]):
                    k = norm(v)
                    if k.startswith("all store"): grp["all"] = j
                    elif k.startswith("own same"): grp["own"] = j
                    elif k.startswith("franchise same"): grp["fran"] = j
                hdr = [norm(v) for v in allrows[i + 1]]
                tt = next((j for j, v in enumerate(hdr) if v.startswith("total sales this")), None)
                tl = next((j for j, v in enumerate(hdr) if v.startswith("total sales last")), None)
                rows_out = []
                for row in allrows[i + 2:]:
                    name = str(row[0]).strip() if row[0] else ""
                    if not name:
                        break
                    rec = {"cat": name.title() if name.lower() != "totals" else "Total"}
                    for g, j in grp.items():
                        rec[g] = r(num(row[j]), 4)
                        rec[g + "_cp"] = r(num(row[j + 2]), 4) if j + 2 < len(row) else None
                    if tt is not None: rec["this"] = r(num(row[tt]), 0)
                    if tl is not None: rec["last"] = r(num(row[tl]), 0)
                    rows_out.append(rec)
                    if name.lower() == "totals":
                        break
                blocks[kind] = rows_out
                i += len(rows_out) + 2
            else:
                i += 1
        if blocks:
            return blocks
    issue("warn", src, "Category growth summary not found.")
    return {}


# --------------------------------------------------------------------------- KPI performance
def month_key(label):
    m = MONTH_RE.match(str(label).strip()) if label else None
    if not m or m.group(1).lower() not in MONTHS:
        return None
    return f"20{m.group(2)}-{MONTHS[m.group(1).lower()]:02d}"


def parse_kpi(path, wb, heads):
    src = os.path.basename(path)
    levels = {}
    # Scoring sheets (*_Wise): the official target/actual/weight used for the score.
    # Some scoring targets differ from Compiled Detail (e.g. stock loss), so they win.
    scoring = {}
    for t, rows in heads.items():
        hi = find_header_row(rows, {"metric", "metric weight", "target", "actual"})
        if hi is None or "head name" in {norm(v) for v in rows[hi]}:
            continue
        hm = header_map(rows[hi])
        smonth = next((month_key(m.group(1)) for v in rows[hi] if v for m in [re.search(r"([A-Za-z]{3}'\d{2})", str(v))] if m and "score" in norm(v)), None)
        mi, wi, ti, ai = hm["metric"], hm["metric weight"], hm["target"], hm["actual"]
        for row in wb[t].iter_rows(min_row=hi + 2, values_only=True):
            if mi >= len(row) or not row[mi]:
                continue
            metric = str(row[mi]).strip()
            key = next((str(v).strip() for v in row if isinstance(v, str) and re.match(r"^\d+", v.strip()) and v.strip().endswith(metric)), None)
            if key:
                scoring[("zonal" if "zonal" in t.lower() else "rho", key)] = (smonth, num(row[ti]), num(row[ai]), num(row[wi]))
    for t, rows in heads.items():
        hi = find_header_row(rows, {"head name", "metric", "calc type"})
        if hi is None:
            continue
        hm = header_map(rows[hi])
        lvl = "zonal" if "zonal" in t.lower() else "rho"
        mcols = [(j, month_key(v)) for j, v in enumerate(rows[hi]) if month_key(v)]
        seen, tcols, acols = set(), [], []
        for j, mk in mcols:
            (acols if mk in seen else tcols).append((j, mk))
            seen.add(mk)
        if not tcols or not acols:
            issue("error", src, f"Sheet '{t}': target/actual month columns not found.")
            continue
        ci, wi = pick(hm, "head name"), pick(hm, "metric weight")
        rowsout = []
        for i, row in enumerate(wb[t].iter_rows(values_only=True)):
            if i <= hi or not row[ci] or not row[hm["metric"]]:
                continue
            metric = str(row[hm["metric"]]).strip()
            head = str(row[ci]).strip()
            w = num(row[wi]) if wi is not None else None
            hid = num(row[hm["id"]]) if "id" in hm else None
            sc = scoring.get((lvl, f"{int(hid)}{metric}")) if hid is not None else None
            if sc and sc[3] is not None:
                w = sc[3]
            rec = {"head": head, "cat": str(row[hm["category"]]).strip(), "metric": metric,
                   "w": w, "dir": norm(row[hm["direction"]]), "calc": norm(row[hm["calc type"]]),
                   "t": {mk: num(row[j]) for j, mk in tcols if j < len(row) and num(row[j]) is not None},
                   "a": {mk: num(row[j]) for j, mk in acols if j < len(row) and num(row[j]) is not None}}
            if sc and sc[0]:
                if sc[1] is not None: rec["t"][sc[0]] = sc[1]
                if sc[2] is not None: rec["a"][sc[0]] = sc[2]
            rowsout.append(rec)
        levels[lvl] = rowsout
    months = sorted({m for rows in levels.values() for rec in rows for m in rec["a"]})
    if not months:
        issue("error", src, "Performance file has no actual values for any month.")
        return None
    # official scores (for validation only)
    official = {}
    for t, rows in heads.items():
        f = {norm(v) for row in rows for v in row if v is not None}
        if {"rho name", "score", "rank"} <= f:
            for row in wb[t].iter_rows(min_row=2, values_only=True):
                if row[0] and num(row[1]) is not None:
                    official[("rho", str(row[0]).strip())] = num(row[1])
                if len(row) > 5 and row[4] and num(row[5]) is not None:
                    official[("zonal", str(row[4]).strip())] = num(row[5])
    return {"file": src, "months": months, "levels": levels, "official": official}


def achievement(t, a, direction):
    if t is None or a is None:
        return None
    if direction.startswith("lower"):
        if a <= 0:
            return 1.0
        v = t / a
    else:
        if t == 0:
            return 1.0 if a >= 0 else 0.0
        v = a / t
    return max(0.0, min(1.0, v))


def month_scores(rows, month):
    """Weighted score per head for one month: sum(ach*w)/sum(w of metrics with data)."""
    acc = {}
    for rec in rows:
        t, a = rec["t"].get(month), rec["a"].get(month)
        ach = achievement(t, a, rec["dir"])
        if ach is None or not rec["w"]:
            continue
        s = acc.setdefault(rec["head"], [0.0, 0.0])
        s[0] += ach * rec["w"]
        s[1] += rec["w"]
    return {h: v[0] / v[1] for h, v in acc.items() if v[1]}


def merge_kpi(files):
    """Each month's figures come from the file that covers the latest month (newest file)."""
    files = sorted(files, key=lambda f: max(f["months"]))
    merged = {}
    for f in files:
        for lvl, rows in f["levels"].items():
            for rec in rows:
                key = (lvl, rec["head"], rec["metric"])
                m = merged.setdefault(key, {k: rec[k] for k in ("head", "cat", "metric", "w", "dir", "calc")} | {"t": {}, "a": {}})
                for mk in f["months"]:
                    if mk in rec["a"]:
                        m["a"][mk] = rec["a"][mk]
                        if mk in rec["t"]:
                            m["t"][mk] = rec["t"][mk]
                m["w"] = rec["w"] or m["w"]
    out = {"rho": [], "zonal": []}
    for (lvl, _, _), m in merged.items():
        m["t"] = {k: r(v, 6) for k, v in m["t"].items()}
        m["a"] = {k: r(v, 6) for k, v in m["a"].items()}
        out[lvl].append(m)
    months = sorted({mk for f in files for mk in f["months"]})
    # validate our scoring against the file's own SCORE sheet
    for f in files:
        if not f["official"]:
            continue
        mk = max(f["months"])
        for lvl in ("rho", "zonal"):
            ours = month_scores(f["levels"].get(lvl, []), mk)
            bad = [h for (l, h), v in f["official"].items() if l == lvl and h in ours and abs(ours[h] - v) > 0.002]
            checked = sum(1 for (l, h) in f["official"] if l == lvl and h in ours)
            if bad:
                issue("warn", f["file"], f"{lvl.upper()} score differs from the file's SCORE sheet for {len(bad)} of {checked} heads (e.g. {bad[0]}).")
            elif checked:
                issue("info", f["file"], f"{lvl.upper()} scores for {mk} match the file's SCORE sheet ({checked} heads checked).")
    return out, months


# --------------------------------------------------------------------------- outlet P&L
def parse_pnl(path, wb, heads):
    src = os.path.basename(path)
    month, sheet = None, None
    for t, rows in heads.items():
        hi = find_header_row(rows, {"code", "sales revenue", "total outlet level opex"})
        if hi is not None:
            sheet = (t, hi)
            for row in rows[:hi]:
                for v in row:
                    m = re.search(r"([A-Za-z]{3})[^A-Za-z0-9]*(\d{4})", str(v)) if isinstance(v, str) and "period" in v.lower() else None
                    if m and m.group(1).lower() in MONTHS:
                        month = f"{m.group(2)}-{MONTHS[m.group(1).lower()]:02d}"
            break
    if not sheet or not month:
        issue("error", src, "P&L file: summary sheet or 'Period' label not found.")
        return None
    t, hi = sheet
    hdr = [norm(v) for v in heads[t][hi]]
    def cols(name):
        return [i for i, h in enumerate(hdr) if h.startswith(name)]
    ci = {"c": cols("code"), "n": cols("outlet name"), "ld": cols("launching date"), "sft": cols("sft"),
          "s": cols("sales revenue"), "gp": cols("gp"), "oi": cols("other income"), "ox": cols("total outlet level opex"),
          "g": cols("outlet level gain/loss"), "ofc": cols("operating financing cost"), "p": cols("outlet level p/(l) after ofc"),
          "ff": cols("foot fall"), "bs": cols("basket size")}
    gp_only = [i for i in ci["gp"] if hdr[i] == "gp"]
    out = []
    for i, row in enumerate(wb[t].iter_rows(values_only=True)):
        if i <= hi:
            continue
        code = str(row[ci["c"][0]] or "").strip()
        if not CODE_RE.match(code):
            continue
        g = lambda k, n=0: row[ci[k][n]] if len(ci[k]) > n and ci[k][n] < len(row) else None
        rec = {"c": code, "n": str(g("n") or "").strip(), "s": num(g("s")), "gp": num(row[gp_only[0]]) if gp_only else None,
               "oi": num(g("oi")), "ox": num(g("ox")), "g": num(g("g")), "ofc": num(g("ofc")), "p": num(g("p")),
               "g0": num(g("g", 1)), "p0": num(g("p", 1)), "sft": num(g("sft")), "ff": num(g("ff")), "bs": num(g("bs"))}
        d = as_date(g("ld"))
        if d:
            rec["ld"] = d.isoformat()
        out.append({k: (r(v, 0) if isinstance(v, float) and k not in ("bs",) else r(v, 2) if isinstance(v, float) else v) for k, v in rec.items() if v not in (None, "")})
    # Transposed detail sheet: line items down column A, one outlet per column
    lines, detail = [], {}
    for t2 in heads:
        top = [list(x) for x in wb[t2].iter_rows(max_row=12, max_col=1, values_only=True)]
        if not any(norm(x[0]) == "code" for x in top if x):
            continue
        grid = [list(x) for x in wb[t2].iter_rows(max_row=80, values_only=True)]
        crow = next(j for j, x in enumerate(grid) if x and norm(x[0]) == "code")
        codes = [str(v).strip() if v else "" for v in grid[crow]]
        labels = [norm(x[0]) for x in grid]
        try:
            a, b = labels.index("gpoi"), labels.index("total outlet level opex")
        except ValueError:
            continue
        keep = [j for j in range(a + 1, b) if grid[j][0]]
        lines = [str(grid[j][0]).strip() for j in keep]
        for col, code in enumerate(codes):
            if col and CODE_RE.match(code) and code not in detail:
                detail[code] = [r(num(grid[j][col]) if col < len(grid[j]) else None, 0) for j in keep]
        break
    if not detail:
        issue("warn", src, "P&L cost breakdown sheet not found; outlet cost details unavailable.")
    loss = sum(1 for o in out if (o.get("p") or 0) < 0)
    issue("info", src, f"P&L {month}: {len(out)} outlets, {loss} loss-making after financing cost.")
    return {"file": src, "month": month, "outlets": out, "lines": lines, "detail": detail}


# --------------------------------------------------------------------------- outlet master
def parse_master(path, wb, heads):
    src = os.path.basename(path)
    best = None
    for t, rows in heads.items():
        hi = find_header_row(rows, {"code", "zonal", "format"})
        if hi is not None:
            best = (t, hi)
            break
    t, hi = best
    hm = header_map(heads[t][hi])
    cols = {"c": ("code",), "n": ("outlet name",), "rl": ("leader",), "rh": ("regional head hr name", "regional head"),
            "zn": ("zonal",), "zh": ("zonal hr name",), "zc": ("zonal contact",), "ld": ("launching date",),
            "sft": ("sft",), "fmt": ("format",), "div": ("division",), "dis": ("district",), "area": ("area",),
            "pnp": ("pnp non pnp status", "pnp"), "own": ("status",), "loc": ("location type",),
            "den": ("population density",), "inc": ("income level",), "geo": ("geo location",)}
    idx = {k: pick(hm, *v) for k, v in cols.items()}
    out, latest = [], None
    for i, row in enumerate(wb[t].iter_rows(values_only=True)):
        if i <= hi:
            continue
        g = lambda k: row[idx[k]] if idx[k] is not None and idx[k] < len(row) else None
        code = str(g("c") or "").strip()
        if not CODE_RE.match(code):
            continue
        rec = {}
        for k in cols:
            v = g(k)
            if k == "ld":
                d = as_date(v)
                v = d.isoformat() if d else None
                if d and (latest is None or d > latest):
                    latest = d
            elif k == "sft":
                v = num(v)
            elif v is not None:
                v = str(v).strip()
                if k == "own":
                    v = "Own" if v.upper().startswith("OWN") else "Franchise" if v.upper().startswith("FR") else v
            if v not in (None, ""):
                rec[k] = v
        out.append(rec)
    if not out:
        issue("error", src, "Outlet master sheet has no outlet rows.")
        return None
    return {"file": src, "sheet": t, "latest_launch": latest, "outlets": out}


# --------------------------------------------------------------------------- main
def month_end(d):
    return d.replace(day=calendar.monthrange(d.year, d.month)[1])


def build(root):
    found = {"business": [], "kpi": [], "master": [], "pnl": []}
    for folder in FOLDERS:
        files = excel_files(os.path.join(root, folder))
        if not files:
            issue("warn", folder, "Folder is empty or could not be read.")
        for p in files:
            kind, payload = classify(p)
            if not kind:
                continue
            wb, heads = payload
            parsed = {"business": parse_business, "kpi": parse_kpi, "master": parse_master, "pnl": parse_pnl}[kind](p, wb, heads)
            wb.close()
            if parsed:
                parsed["folder"] = folder
                found[kind].append(parsed)
                print(f"recognised {kind}: {os.path.basename(p)} ({folder})")

    biz = found["business"]
    td_pool = [b for b in biz if b["folder"] == "tilldate"] or biz
    if not td_pool:
        issue("error", "tilldate", "No till-date business report found.")
        return None
    till = max(td_pool, key=lambda b: b["date"])
    expected = till["date"].replace(day=1) - dt.timedelta(days=1)
    closed = [b for b in biz if b["date"] == month_end(b["date"]) and b is not till]
    me = next((b for b in closed if b["date"] == expected), None)
    if me is None and closed:
        me = max(closed, key=lambda b: b["date"])
        issue("warn", "monthend", f"{expected:%B %Y} month-end file not uploaded yet; showing {me['date']:%B %Y}.")
    elif me is None:
        issue("warn", "monthend", "No closed month-end file found.")

    master = max(found["master"], key=lambda m: (m["latest_launch"] or dt.date.min, len(m["outlets"])), default=None)
    if master is None:
        issue("warn", "performance", "No outlet master (zone distribution) found; filters will use the names in the sales files.")

    kpi, kpi_months = ({"rho": [], "zonal": []}, [])
    if found["kpi"]:
        kpi, kpi_months = merge_kpi(found["kpi"])
    else:
        issue("warn", "performance", "No KPI performance file found.")

    pnl = None
    if found["pnl"]:
        bym = {}
        for f in found["pnl"]:
            if f["month"] not in bym or len(f["outlets"]) > len(bym[f["month"]]["outlets"]):
                bym[f["month"]] = f
        months = sorted(bym)[-12:]
        lines = bym[months[-1]]["lines"]
        detail = {}
        for m in months:  # align every month's cost lines to the latest month's line list
            f = bym[m]
            idx = [f["lines"].index(l) if l in f["lines"] else None for l in lines]
            detail[m] = {c: [v[i] if i is not None else None for i in idx] for c, v in f["detail"].items()}
        pnl = {"months": months, "summary": {m: bym[m]["outlets"] for m in months}, "lines": lines, "detail": detail}
    else:
        issue("warn", "performance", "No outlet P&L file found; the Loss-making outlets page is empty.")

    # join checks
    if master:
        mcodes = {o["c"] for o in master["outlets"]}
        for rep, name in ((till, "Till-date"), (me, "Month-end")):
            if rep:
                miss = sorted(o["c"] for o in rep["outlets"] if o["c"] not in mcodes and o.get("s"))
                if miss:
                    issue("info", rep["file"], f"{name}: {len(miss)} outlets with sales are not in the outlet master (shown as Unmapped): {', '.join(miss[:12])}{'…' if len(miss) > 12 else ''}")

    def rep_out(b):
        if not b:
            return None
        d = b["date"]
        return {"file": b["file"], "date": d.isoformat(), "day": d.day, "dim": calendar.monthrange(d.year, d.month)[1],
                "closed": d == month_end(d), "sply_label": b.get("sply_label"), "splm_label": b.get("splm_label"),
                "categories": b.get("categories", {}), "outlets": b["outlets"]}

    return {
        "generated": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
        "tilldate": rep_out(till),
        "monthend": rep_out(me),
        "master": {"file": master["file"], "outlets": master["outlets"]} if master else None,
        "pnl": pnl,
        "kpi": {"months": kpi_months, "files": [f["file"] for f in found["kpi"]], **kpi},
        "sources": [{"folder": b["folder"], "file": b["file"], "type": "Business report", "date": b["date"].isoformat()} for b in biz]
                   + [{"folder": f["folder"], "file": f["file"], "type": "KPI performance", "date": max(f["months"])} for f in found["kpi"]]
                   + [{"folder": f["folder"], "file": f["file"], "type": "Outlet P&L", "date": f["month"]} for f in found["pnl"]]
                   + [{"folder": m["folder"], "file": m["file"], "type": "Outlet master", "date": (m["latest_launch"] or dt.date.min).isoformat()} for m in found["master"]],
        "issues": ISSUES,
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--local", help="folder with tilldate/ monthend/ performance/ subfolders")
    args = ap.parse_args()
    tmp = None
    if args.local:
        root = args.local
    else:
        tmp = root = tempfile.mkdtemp()
        download(root)
    try:
        data = build(root)
    finally:
        if tmp:
            shutil.rmtree(tmp, ignore_errors=True)
    if not data or any(i["level"] == "error" and i["source"] in ("tilldate",) for i in ISSUES):
        print("Critical problem: data.json NOT updated (the dashboard keeps the last good data).")
        sys.exit(1)
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    new = json.loads(json.dumps(data, default=str))
    try:
        with open(OUT, encoding="utf-8") as fh:
            old = json.load(fh)
        if {k: v for k, v in old.items() if k != "generated"} == {k: v for k, v in new.items() if k != "generated"}:
            print("No data changes; data.json left as is.")
            return
    except (OSError, ValueError):
        pass
    tmp_out = OUT + ".tmp"
    with open(tmp_out, "w", encoding="utf-8") as fh:
        json.dump(data, fh, ensure_ascii=False, separators=(",", ":"), default=str)
    os.replace(tmp_out, OUT)  # atomic: visitors never see a half-written file
    print(f"Wrote {OUT} ({os.path.getsize(OUT)/1024:.0f} KB)")


if __name__ == "__main__":
    main()
