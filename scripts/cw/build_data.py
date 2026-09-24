#!/usr/bin/env python3
"""Normalize Shwapno consumable, wastage, sales, hierarchy and target files."""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

try:
    import xlrd
except ImportError:  # optional at import time; required only for true legacy .xls files
    xlrd = None


OUTLET_CODE_RE = re.compile(r"^[A-Z][0-9]{3,4}$")


def clean_text(value: Any) -> str:
    return " ".join(str(value or "").replace("\u00a0", " ").split()).strip()


def normalize_outlet_code(value: Any) -> str:
    value = clean_text(value).upper()
    return value if OUTLET_CODE_RE.fullmatch(value) else ""


def parse_date(value: Any) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = clean_text(value)
    for pattern in ("%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return dt.datetime.strptime(text, pattern).date()
        except ValueError:
            pass
    return None


def iso_date(value: dt.date | None) -> str | None:
    return value.isoformat() if value else None


def parse_number(value: Any) -> float:
    if value in (None, ""):
        return 0.0
    text = clean_text(value).replace(",", "")
    if text.startswith("(") and text.endswith(")"):
        text = f"-{text[1:-1]}"
    try:
        return float(text)
    except ValueError:
        return 0.0


# Benchmark targets are lower-is-better control limits expressed in percent.
# Anything outside this band means the Target.txt column was entered in the
# wrong unit, so the build fails loudly instead of publishing a 100x error.
TARGET_MIN = 0.0001  # 0.01%
TARGET_MAX = 0.5000  # 50.00%


def parse_percent(value: Any) -> float | None:
    """Parse a benchmark target column that is documented in percent.

    "0.50%" and a bare "0.50" both mean 0.50 percent. The previous
    ``number / 100 if number > 1 else number`` heuristic silently read a bare
    ``0.5`` as 50%, so the unit is now fixed by the column, not guessed.
    """
    text = clean_text(value).replace(",", "")
    if not text:
        return None
    if text.endswith("%"):
        text = text[:-1].strip()
    try:
        return float(text) / 100
    except ValueError:
        return None


def canonical_criteria(format_name: Any, pnp_status: Any, ownership: Any) -> str:
    pnp = clean_text(pnp_status).upper().replace(" ", "-")
    pnp = "Non-PNP" if pnp in {"NON-PNP", "NONPNP"} else "PNP"
    owner = clean_text(ownership).upper()
    owner = "Own" if owner in {"OWN", "OWNED"} else "FR" if owner == "FR" else clean_text(ownership)
    return f"{clean_text(format_name)}-{pnp}-{owner}"


def first_present(values: Iterable[str]) -> str:
    for value in values:
        if clean_text(value):
            return clean_text(value)
    return ""


TARGET_COLUMNS = {
    "consumableTarget": "Consumable % On Sales",
    "wastageSalesTarget": "Wastage % On Sales",
    "wastagePnpTarget": "Wastage % On PNP Sales",
}


MONTH_NUMBERS = {
    "jan": 1, "january": 1,
    "feb": 2, "february": 2,
    "mar": 3, "march": 3,
    "apr": 4, "april": 4,
    "may": 5,
    "jun": 6, "june": 6,
    "jul": 7, "july": 7,
    "aug": 8, "august": 8,
    "sep": 9, "sept": 9, "september": 9,
    "oct": 10, "october": 10,
    "nov": 11, "november": 11,
    "dec": 12, "december": 12,
}


def source_file_kind(path: Path) -> str | None:
    """Classify a downloaded source file from its filename.

    Drive refreshes often include dates or Windows copy suffixes, so filenames
    do not need to be exact. Matching stays deliberately narrow to avoid using
    an unrelated workbook from the shared folder.
    """
    if not path.is_file():
        return None
    suffix = path.suffix.casefold()
    label = re.sub(r"[^a-z0-9]+", " ", path.stem.casefold()).strip()
    words = set(label.split())

    if suffix == ".txt" and "target" in words:
        return "targets"
    if suffix in {".xlsx", ".xlsm"} and {"sales", "till"} <= words:
        return "sales"
    if suffix in {".xlsx", ".xlsm"} and {"zone", "distribution"} <= words:
        return "zones"
    sap_suffixes = {".xls", ".xlsx", ".xlsm", ".txt", ".csv"}
    if suffix in sap_suffixes and "consumable" in words:
        return "consumable"
    if suffix in sap_suffixes and "wastage" in words:
        return "wastage"
    return None


def filename_freshness(path: Path) -> tuple[int, int, int, str]:
    """Return a deterministic freshness score derived only from the filename."""
    stem = path.stem.casefold()
    best_ordinal = 0
    date_precision = 0

    # Full numeric dates: DD.MM.YYYY / DD-MM-YYYY / YYYY-MM-DD and variants.
    for match in re.finditer(r"(?<!\d)(\d{1,4})[._ -](\d{1,2})[._ -](\d{1,4})(?!\d)", stem):
        a, b, c = (int(value) for value in match.groups())
        candidates: list[dt.date] = []
        try:
            if a >= 1900:
                candidates.append(dt.date(a, b, c))
            elif c >= 1900:
                candidates.append(dt.date(c, b, a))
        except ValueError:
            pass
        if candidates:
            best_ordinal = max(best_ordinal, max(value.toordinal() for value in candidates))
            date_precision = max(date_precision, 2)

    # Month-name + year, e.g. "Zone Distribution Aug 2026".
    month_pattern = "|".join(sorted(MONTH_NUMBERS, key=len, reverse=True))
    for match in re.finditer(rf"\b({month_pattern})[ '\\._-]*(20\d{{2}})\b", stem):
        month = MONTH_NUMBERS[match.group(1)]
        year = int(match.group(2))
        best_ordinal = max(best_ordinal, dt.date(year, month, 1).toordinal())
        date_precision = max(date_precision, 1)

    # Windows/browser duplicate suffixes, e.g. Target(1), file (2).
    copy_numbers = [int(value) for value in re.findall(r"\((\d+)\)\s*$", stem)]
    version = max(copy_numbers) if copy_numbers else 0
    return (date_precision, best_ordinal, version, path.name.casefold())


def resolve_source_files(input_dir: Path) -> dict[str, Path]:
    """Find the five required dashboard inputs anywhere under input_dir."""
    if not input_dir.is_dir():
        raise FileNotFoundError(f"Source directory not found: {input_dir}")

    grouped: dict[str, list[Path]] = {
        "targets": [], "sales": [], "zones": [], "consumable": [], "wastage": []
    }
    for path in input_dir.rglob("*"):
        kind = source_file_kind(path)
        if kind:
            grouped[kind].append(path)

    missing = [kind for kind, candidates in grouped.items() if not candidates]
    if missing:
        readable = {
            "targets": "Target*.txt",
            "sales": "Sales-Till*.xlsx",
            "zones": "Zone Distribution*.xlsx",
            "consumable": "CONSUMABLE*.xls",
            "wastage": "WASTAGE*.xls",
        }
        expected = ", ".join(readable[kind] for kind in missing)
        found = ", ".join(sorted(path.name for path in input_dir.rglob("*") if path.is_file()))
        raise FileNotFoundError(
            f"Missing required Drive source type(s): {', '.join(missing)}. "
            f"Expected names like: {expected}. Found: {found or 'no files'}."
        )

    selected: dict[str, Path] = {}
    for kind, candidates in grouped.items():
        selected[kind] = max(candidates, key=filename_freshness)
    return selected


def read_targets(path: Path) -> tuple[dict[str, dict[str, Any]], list[dict[str, Any]]]:
    targets: dict[str, dict[str, Any]] = {}
    ordered: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        headers = {clean_text(name) for name in (reader.fieldnames or [])}
        required = {"Final Criteria", *TARGET_COLUMNS.values()}
        missing = sorted(required - headers)
        if missing:
            raise ValueError(
                f"{path.name} is missing columns: {', '.join(missing)}. "
                f"Found: {', '.join(sorted(headers))}."
            )
        for line_number, row in enumerate(reader, start=2):
            criteria = clean_text(row.get("Final Criteria"))
            if not criteria:
                continue
            item: dict[str, Any] = {"criteria": criteria}
            for field, column in TARGET_COLUMNS.items():
                raw = row.get(column)
                parsed = parse_percent(raw)
                if parsed is None:
                    raise ValueError(
                        f"{path.name} line {line_number}: '{column}' for "
                        f"'{criteria}' is not a percentage (found {raw!r})."
                    )
                if not TARGET_MIN <= parsed <= TARGET_MAX:
                    raise ValueError(
                        f"{path.name} line {line_number}: '{column}' for "
                        f"'{criteria}' resolves to {parsed * 100:.4f}%, outside the "
                        f"expected {TARGET_MIN * 100:g}%-{TARGET_MAX * 100:g}% band. "
                        f"Enter targets in percent, for example 0.50%."
                    )
                item[field] = parsed
            if criteria.casefold() in targets:
                raise ValueError(f"{path.name}: duplicate Final Criteria '{criteria}'.")
            targets[criteria.casefold()] = item
            ordered.append(item)
    if not ordered:
        raise ValueError(f"{path.name} contains no benchmark criteria rows.")
    return targets, ordered


def read_zone_distribution(path: Path, targets: dict[str, dict[str, Any]]) -> tuple[dict[str, dict[str, Any]], dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet_name = "Final_Zone Dis" if "Final_Zone Dis" in workbook.sheetnames else workbook.sheetnames[0]
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    headers = [clean_text(value) for value in next(rows)]
    index = {header: position for position, header in enumerate(headers)}
    required = {"CODE", "Outlet Name", "Format", "Division", "District", "PNP Non PNP status", "Status"}
    missing = sorted(required - set(index))
    if missing:
        raise ValueError(f"Zone distribution is missing columns: {', '.join(missing)}")

    def value(row: tuple[Any, ...], name: str) -> Any:
        position = index.get(name)
        return row[position] if position is not None and position < len(row) else None

    records: dict[str, dict[str, Any]] = {}
    duplicates: list[str] = []
    unknown_criteria: Counter[str] = Counter()
    row_count = 0
    for row in rows:
        code = normalize_outlet_code(value(row, "CODE"))
        if not code:
            continue
        row_count += 1
        if code in records:
            duplicates.append(code)
        criteria = canonical_criteria(
            value(row, "Format"), value(row, "PNP Non PNP status"), value(row, "Status")
        )
        target = targets.get(criteria.casefold())
        if target is None:
            unknown_criteria[criteria] += 1
        launch_date = parse_date(value(row, "Launching Date"))
        records[code] = {
            "code": code,
            "name": clean_text(value(row, "Outlet Name")) or code,
            "regionalLeader": first_present(
                [value(row, "Leader"), value(row, "Regional Head HR Name")]
            ),
            "regionalLeaderFullName": clean_text(value(row, "Regional Head HR Name")),
            "zone": first_present([value(row, "Zonal"), value(row, "Zonal HR Name")]),
            "zonalFullName": clean_text(value(row, "Zonal HR Name")),
            "division": clean_text(value(row, "Division")),
            "district": clean_text(value(row, "District")),
            "area": clean_text(value(row, "Area")),
            "format": clean_text(value(row, "Format")),
            "pnpStatus": "PNP" if "NON" not in clean_text(value(row, "PNP Non PNP status")).upper() else "Non-PNP",
            "ownership": "Own" if clean_text(value(row, "Status")).upper() == "OWN" else clean_text(value(row, "Status")),
            "locationType": clean_text(value(row, "Location Type")),
            "locationClass": clean_text(value(row, "Location Type(Dv,Ds,T)")),
            "populationDensity": clean_text(value(row, "Population Density")),
            "incomeLevel": clean_text(value(row, "Income level")),
            "launchDate": iso_date(launch_date),
            "sft": parse_number(value(row, "SFT")) or None,
            "criteria": target["criteria"] if target else criteria,
            "consumableTarget": target["consumableTarget"] if target else None,
            "wastageSalesTarget": target["wastageSalesTarget"] if target else None,
            "wastagePnpTarget": target["wastagePnpTarget"] if target else None,
            "mapped": True,
        }
    workbook.close()
    quality = {
        "sheet": sheet_name,
        "rows": row_count,
        "duplicateCodes": sorted(set(duplicates)),
        "unknownCriteria": dict(unknown_criteria),
    }
    return records, quality


def read_sales(path: Path) -> tuple[dict[tuple[str, dt.date], dict[str, float]], dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet_name = "Comparative" if "Comparative" in workbook.sheetnames else workbook.sheetnames[0]
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    headers = [clean_text(value) for value in next(rows)]
    index = {header: position for position, header in enumerate(headers)}
    required = {"Outlet Code", "Date", "Division", "POS NSI"}
    missing = sorted(required - set(index))
    if missing:
        raise ValueError(f"Sales file is missing columns: {', '.join(missing)}")

    def value(row: tuple[Any, ...], name: str) -> Any:
        position = index.get(name)
        return row[position] if position is not None and position < len(row) else None

    daily: dict[tuple[str, dt.date], dict[str, float]] = defaultdict(
        lambda: {"sales": 0.0, "pnpSales": 0.0}
    )
    division_rows: Counter[str] = Counter()
    division_value: dict[str, float] = defaultdict(float)
    dates: list[dt.date] = []
    invalid_rows = 0
    row_count = 0
    detail_total = 0.0
    reported_total: float | None = None
    for row in rows:
        raw_code = clean_text(value(row, "Outlet Code"))
        sales_value = parse_number(value(row, "POS NSI"))
        # The export ends with a "Total (999)" footer row. Capture it so the
        # detail rows can be reconciled against the source's own total.
        if raw_code.upper().startswith("TOTAL"):
            reported_total = sales_value
            continue
        code = normalize_outlet_code(raw_code)
        date = parse_date(value(row, "Date"))
        if not code or date is None:
            if raw_code or sales_value:
                invalid_rows += 1
            continue
        division = clean_text(value(row, "Division"))
        daily[(code, date)]["sales"] += sales_value
        # User-confirmed business rule: FRESH PRODUCE is PNP Sales.
        if division.upper() == "FRESH PRODUCE":
            daily[(code, date)]["pnpSales"] += sales_value
        division_rows[division] += 1
        division_value[division] += sales_value
        detail_total += sales_value
        dates.append(date)
        row_count += 1
    workbook.close()

    reconciliation: dict[str, Any] = {
        "detailTotal": round_money(detail_total),
        "reportedTotal": None if reported_total is None else round_money(reported_total),
        "difference": None,
        "differenceRatio": None,
        "matches": None,
    }
    if reported_total is not None:
        difference = detail_total - reported_total
        reconciliation["difference"] = round_money(difference)
        reconciliation["differenceRatio"] = (
            abs(difference) / abs(reported_total) if reported_total else None
        )
        # Tight enough to surface a real export gap, loose enough to ignore
        # currency rounding. Anything above 0.1% blocks the build in
        # validate_payload; below that it is reported, not fatal.
        reconciliation["matches"] = abs(difference) <= max(1.0, abs(reported_total) * 1e-7)

    quality = {
        "sheet": sheet_name,
        "rows": row_count,
        "invalidRows": invalid_rows,
        "dateMin": iso_date(min(dates) if dates else None),
        "dateMax": iso_date(max(dates) if dates else None),
        "outlets": len({code for code, _ in daily}),
        "divisionRows": dict(sorted(division_rows.items())),
        "divisionValue": {name: round_money(total) for name, total in sorted(division_value.items())},
        "reconciliation": reconciliation,
    }
    return daily, quality


# SAP ALV exports can arrive in several technically different formats even when
# SAP gives them the same .xls extension: UTF-16 tab text, pipe-delimited text,
# a real legacy Excel workbook, or a modern Excel workbook.  Keep the parser
# format-tolerant while still locating columns by business labels so a layout
# change cannot silently shift the data into the wrong fields.
SAP_COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "postingDate": (
        "Pstng Date", "Posting Date", "Pstg Date", "Postg Date", "Post Date",
        "Posting Dt", "Pstng Dt",
    ),
    "plant": ("Plnt", "Plant", "Plant Code"),
    "amount": (
        "Amount in LC", "Amount in local currency", "Amount LC", "Amt in LC",
        "Amt.in loc.cur.", "Amt in loc cur", "Local Currency Amount",
    ),
    "name": ("Name 1", "Name1", "Name", "Plant Name"),
    "material": ("Material", "Material No", "Material Number"),
    "materialDescription": (
        "Material Description", "Material description", "Description",
        "Material Desc", "Mat. Description",
    ),
    "quantity": (
        "Quantity in UnE", "Qty in UnE", "Quantity", "Qty",
        "Qty in unit of entry", "Quantity in unit of entry",
    ),
    "unit": ("EUn", "BUn", "UoM", "Unit", "Unit of Entry"),
    "movement": ("MvT", "Movement type", "Movement Type", "MvT.", "Mvt", "Mvmt Type"),
    "documentDate": ("Doc. Date", "Doc Date", "Document Date", "Document Dt"),
}
SAP_REQUIRED_COLUMNS = ("postingDate", "plant", "amount")
SAP_HEADER_SEARCH_ROWS = 3000


def normalize_sap_label(value: Any) -> str:
    """Normalize punctuation/spacing differences in SAP ALV column labels."""
    return re.sub(r"[^a-z0-9]+", " ", clean_text(value).casefold()).strip()


SAP_NORMALIZED_ALIASES = {
    field: {normalize_sap_label(alias) for alias in aliases}
    for field, aliases in SAP_COLUMN_ALIASES.items()
}


def locate_sap_header(cells: list[str]) -> dict[str, int] | None:
    """Return a field -> column index map if this row is an ALV header."""
    if len(cells) < 3:
        return None
    lookup = {
        normalize_sap_label(cell): position
        for position, cell in enumerate(cells)
        if normalize_sap_label(cell)
    }
    mapping: dict[str, int] = {}
    for field, aliases in SAP_NORMALIZED_ALIASES.items():
        for alias in aliases:
            position = lookup.get(alias)
            if position is not None:
                mapping[field] = position
                break
    if any(field not in mapping for field in SAP_REQUIRED_COLUMNS):
        return None
    return mapping


def _detect_text_encoding(sample: bytes) -> str:
    if sample.startswith((b"\xff\xfe", b"\xfe\xff")):
        return "utf-16"
    if sample.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"
    # SAP Unicode exports sometimes lose their BOM during upload/copying.
    if sample:
        nul_ratio = sample.count(b"\x00") / len(sample)
        if nul_ratio > 0.20:
            even_nuls = sample[0::2].count(b"\x00")
            odd_nuls = sample[1::2].count(b"\x00")
            return "utf-16-be" if even_nuls > odd_nuls else "utf-16-le"
    try:
        sample.decode("utf-8")
        return "utf-8-sig"
    except UnicodeDecodeError:
        return "cp1252"


def _detect_text_delimiter(sample_text: str) -> str:
    lines = [line for line in sample_text.splitlines()[:80] if line.strip()]
    if not lines:
        return "\t"
    candidates = ("\t", "|", ";", ",")
    scores = {
        delimiter: sum(line.count(delimiter) for line in lines)
        for delimiter in candidates
    }
    delimiter = max(scores, key=scores.get)
    return delimiter if scores[delimiter] > 0 else "\t"


def _iter_text_rows(path: Path):
    sample = path.read_bytes()[:131072]
    encoding = _detect_text_encoding(sample)
    preview = sample.decode(encoding, errors="replace")
    delimiter = _detect_text_delimiter(preview)
    with path.open("r", encoding=encoding, errors="replace", newline="") as handle:
        for row in csv.reader(handle, delimiter=delimiter):
            # Pipe exports usually carry a decorative empty cell on both edges.
            if delimiter == "|":
                while row and not clean_text(row[0]):
                    row = row[1:]
                while row and not clean_text(row[-1]):
                    row = row[:-1]
            yield row


def _select_xlsx_sheet(path: Path):
    # Loading through BytesIO bypasses openpyxl's filename-extension check, which
    # also lets us read a ZIP/XLSX workbook that SAP happened to name *.xls.
    workbook = load_workbook(io.BytesIO(path.read_bytes()), read_only=True, data_only=True)
    selected = None
    for sheet in workbook.worksheets:
        for line_number, row in enumerate(sheet.iter_rows(values_only=True)):
            if line_number >= SAP_HEADER_SEARCH_ROWS:
                break
            cells = [clean_text(value) for value in row]
            if locate_sap_header(cells) is not None:
                selected = sheet.title
                break
        if selected:
            break
    if not selected:
        workbook.close()
        raise ValueError(f"{path.name}: no SAP header found in any Excel worksheet.")
    sheet = workbook[selected]
    return workbook, sheet


def _iter_xlsx_rows(path: Path):
    workbook, sheet = _select_xlsx_sheet(path)
    try:
        for row in sheet.iter_rows(values_only=True):
            yield list(row)
    finally:
        workbook.close()


def _select_xls_sheet(path: Path):
    if xlrd is None:
        raise RuntimeError(
            f"{path.name}: this is a real legacy Excel .xls workbook. "
            "Install xlrd (already included in the corrected requirements.txt)."
        )
    workbook = xlrd.open_workbook(file_contents=path.read_bytes(), on_demand=True)
    selected_index = None
    for index in range(workbook.nsheets):
        sheet = workbook.sheet_by_index(index)
        for row_number in range(min(sheet.nrows, SAP_HEADER_SEARCH_ROWS)):
            cells = [clean_text(value) for value in sheet.row_values(row_number)]
            if locate_sap_header(cells) is not None:
                selected_index = index
                break
        if selected_index is not None:
            break
    if selected_index is None:
        workbook.release_resources()
        raise ValueError(f"{path.name}: no SAP header found in any legacy Excel worksheet.")
    return workbook, workbook.sheet_by_index(selected_index)


def _iter_xls_rows(path: Path):
    workbook, sheet = _select_xls_sheet(path)
    try:
        for row_number in range(sheet.nrows):
            values: list[Any] = []
            for cell in sheet.row(row_number):
                value = cell.value
                if xlrd is not None and cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        value = xlrd.xldate.xldate_as_datetime(value, workbook.datemode)
                    except Exception:
                        pass
                values.append(value)
            yield values
    finally:
        workbook.release_resources()


def iter_sap_rows(path: Path):
    """Yield rows from text, true .xls, or XLSX-style SAP exports."""
    sample = path.read_bytes()[:16]
    # OLE Compound File = true old .xls.
    if sample.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
        yield from _iter_xls_rows(path)
        return
    # PK ZIP = .xlsx/.xlsm, including incorrectly named .xls files.
    if sample.startswith(b"PK\x03\x04"):
        yield from _iter_xlsx_rows(path)
        return
    yield from _iter_text_rows(path)


def read_sap_text_export(
    path: Path,
) -> tuple[dict[tuple[str, dt.date], float], dict[str, str], dict[str, Any], dict[str, Any]]:
    """Read SAP movement export regardless of text/Excel container format."""
    daily: dict[tuple[str, dt.date], float] = defaultdict(float)
    outlet_names: dict[str, str] = {}
    material_labels: dict[str, tuple[str, str]] = {}
    material_daily: dict[tuple[str, dt.date, str], list[float]] = defaultdict(lambda: [0.0, 0.0])
    movement_rows: Counter[str] = Counter()
    movement_value: dict[str, float] = defaultdict(float)
    material_codes: set[str] = set()
    dates: list[dt.date] = []
    positive_amount_rows = 0
    negative_amount_rows = 0
    row_count = 0
    skipped_rows = 0
    repeated_headers = 0
    columns: dict[str, int] | None = None
    header_row = None
    preview_rows: list[str] = []

    for line_number, row in enumerate(iter_sap_rows(path)):
        cells = [clean_text(cell) for cell in row]
        if len(preview_rows) < 8 and any(cells):
            preview_rows.append(" | ".join(cells[:12])[:500])

        if columns is None:
            if line_number >= SAP_HEADER_SEARCH_ROWS:
                break
            found = locate_sap_header(cells)
            if found is not None:
                columns = found
                header_row = line_number + 1
            continue

        width = max(columns.values()) + 1
        if len(cells) < width:
            continue

        def field(name: str) -> str:
            position = columns.get(name) if columns else None
            return cells[position] if position is not None else ""

        code = normalize_outlet_code(field("plant"))
        posting_date = parse_date(field("postingDate"))
        if not code or posting_date is None:
            repeated = locate_sap_header(cells) is not None
            if any(cells) and not repeated:
                skipped_rows += 1
            else:
                repeated_headers += int(repeated)
            continue

        amount = parse_number(field("amount"))
        movement = field("movement")
        daily[(code, posting_date)] += -amount
        movement_value[movement] += -amount
        name = field("name")
        if name:
            outlet_names[code] = name
        movement_rows[movement] += 1

        material = field("material")
        material_codes.add(material)
        if material not in material_labels:
            material_labels[material] = (field("materialDescription"), field("unit"))
        # Retain every item, posting date and signed movement. A material that
        # is small overall can lead a filtered selection; reversals can change
        # its position too. Blank material codes still contribute to net spend.
        bucket = material_daily[(code, posting_date, material)]
        bucket[0] += -amount
        bucket[1] += -parse_number(field("quantity"))

        dates.append(posting_date)
        row_count += 1
        positive_amount_rows += int(amount > 0)
        negative_amount_rows += int(amount < 0)

    if columns is None:
        sample = " || ".join(preview_rows[:4]) or "<no readable rows>"
        raise ValueError(
            f"{path.name}: could not find the SAP column header within the first "
            f"{SAP_HEADER_SEARCH_ROWS} rows. Required business columns are Posting Date, "
            f"Plant and Amount in LC (common SAP label variants are supported). "
            f"First readable rows: {sample}"
        )
    if row_count == 0:
        raise ValueError(
            f"{path.name}: the column header was found on row {header_row} but no "
            f"movement rows parsed. Check that the Plant column holds outlet codes such as F123/D084."
        )

    magic = path.read_bytes()[:8]
    if magic.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
        detected_format = "xls"
    elif magic.startswith(b"PK\x03\x04"):
        detected_format = "xlsx"
    else:
        detected_format = "text"

    quality = {
        "rows": row_count,
        "headerRow": header_row,
        "detectedFormat": detected_format,
        "columns": {name: position for name, position in sorted(columns.items())},
        "skippedRows": skipped_rows,
        "repeatedHeaderRows": repeated_headers,
        "dateMin": iso_date(min(dates) if dates else None),
        "dateMax": iso_date(max(dates) if dates else None),
        "outlets": len({code for code, _ in daily}),
        "materials": len(material_codes),
        "movementRows": dict(sorted(movement_rows.items())),
        "movementValue": {name: round_money(value) for name, value in sorted(movement_value.items())},
        "issueRows": negative_amount_rows,
        "reversalRows": positive_amount_rows,
    }
    detail = {
        "labels": material_labels,
        "daily": material_daily,
    }
    return daily, outlet_names, quality, detail


def build_material_view(
    consumable_detail: dict[str, Any],
    wastage_detail: dict[str, Any],
) -> tuple[list[list[Any]], dict[str, list[list[Any]]]]:
    """Keep complete outlet/date material totals in a compact shared catalog.

    Ranking happens in the browser after global filters, never at build time.
    Rows are [postingDate, outletCode, materialId, netValue, netQuantity].
    """
    catalog_index: dict[str, int] = {}
    catalog: list[list[Any]] = []

    def material_id(code: str, source: dict[str, Any]) -> int:
        if code not in catalog_index:
            description, unit = source["labels"].get(code, ("", ""))
            catalog_index[code] = len(catalog)
            catalog.append([code, description or code or "Unknown material", unit])
        return catalog_index[code]

    daily = {}
    for key, source in (("consumable", consumable_detail), ("wastage", wastage_detail)):
        daily[key] = [
            [iso_date(date), code, material_id(material, source), round_money(values[0]), round(values[1], 2)]
            for (code, date, material), values in sorted(source["daily"].items())
        ]
    return catalog, daily


def build_calendar(dates: list[str]) -> list[list[Any]]:
    """Date, ISO week and weekday index, so week grouping needs no JS date math."""
    rows: list[list[Any]] = []
    for value in dates:
        parsed = parse_date(value)
        if not parsed:
            continue
        iso = parsed.isocalendar()
        rows.append([value, iso.week, parsed.isoweekday()])
    return rows


def round_money(value: float) -> float:
    return round(value + 0.0, 2)


def build_period_alignment(sources: dict[str, dict[str, Any]]) -> dict[str, Any]:
    """Compare the covered period of every source file.

    Sales is the denominator of every rate. If the consumable or wastage export
    stops earlier than sales, the rate is divided by more days than it earns,
    which understates it. This is reported rather than silently assumed to pass.
    """
    spans = {
        name: (parse_date(quality.get("dateMin")), parse_date(quality.get("dateMax")))
        for name, quality in sources.items()
    }
    starts = [start for start, _ in spans.values() if start]
    ends = [end for _, end in spans.values() if end]
    common_start = max(starts) if starts else None
    common_end = min(ends) if ends else None
    reference_end = spans.get("sales", (None, None))[1]

    details: dict[str, Any] = {}
    for name, (start, end) in spans.items():
        details[name] = {
            "dateMin": iso_date(start),
            "dateMax": iso_date(end),
            "startLagDays": (start - common_start).days if start and common_start else None,
            "endLagDays": (reference_end - end).days if end and reference_end else None,
        }

    aligned = (
        len(starts) == len(spans)
        and len(ends) == len(spans)
        and len(set(starts)) == 1
        and len(set(ends)) == 1
    )
    lagging = sorted(
        name
        for name, detail in details.items()
        if detail["endLagDays"] not in (None, 0) and detail["endLagDays"] > 0
    )
    return {
        "aligned": aligned,
        "commonStart": iso_date(common_start),
        "commonEnd": iso_date(common_end),
        "laggingSources": lagging,
        "sources": details,
    }


def consumable_posting_cutoffs(
    consumable_daily: dict[tuple[str, dt.date], float],
) -> dict[str, str]:
    """The last date each outlet had a consumable movement posted.

    Consumable issues reach SAP in batches, so an outlet's postings stop on the
    day of its last batch while its sales keep running to the end of the file.
    Dividing that outlet's consumable value by the full sales period charges it
    for days it never booked against, which understates the rate.

    A posting day is one that carries consumable value. A day whose issues and
    reversals cancel to exactly zero booked nothing, so it does not extend the
    window; extending it would add sales days against no value and understate
    the rate again. The dashboard applies the same test to the same daily
    values, so the two sides cannot drift apart.

    Returned as ISO strings, which compare correctly in both languages.
    """
    cutoffs: dict[str, dt.date] = {}
    for (code, date), value in consumable_daily.items():
        if not value:
            continue
        current = cutoffs.get(code)
        if current is None or date > current:
            cutoffs[code] = date
    return {code: date.isoformat() for code, date in cutoffs.items()}


def consumable_sales_base(daily_rows: list[dict[str, Any]], cutoffs: dict[str, str]) -> float:
    """Sales that fall on or before each outlet's own consumable cut-off.

    This is the denominator of Consumable % on Sales. An outlet with no
    consumable posting at all contributes nothing: its rate is unknown, not
    zero, so it drops out of the weighted result rather than diluting it.
    ISO dates compare correctly as strings, which is also how the dashboard
    does it, so both sides stay in step.
    """
    total = 0.0
    for row in daily_rows:
        cutoff = cutoffs.get(row["code"])
        if cutoff is not None and row["date"] <= cutoff:
            total += row["sales"]
    return total


def build_consumable_window(
    daily_rows: list[dict[str, Any]],
    cutoffs: dict[str, str],
    activity_codes: set[str],
    sales_last_date: str | None,
    consumable_last_date: str | None,
) -> dict[str, Any]:
    """Diagnostics for the per-outlet consumable denominator."""
    base = consumable_sales_base(daily_rows, cutoffs)
    sales_total = sum(row["sales"] for row in daily_rows)
    consumable_total = sum(row["consumable"] for row in daily_rows)
    truncated = sorted(
        code
        for code, cutoff in cutoffs.items()
        if sales_last_date is not None and cutoff < sales_last_date
    )
    lags = []
    if sales_last_date:
        reference = parse_date(sales_last_date)
        for cutoff in cutoffs.values():
            parsed = parse_date(cutoff)
            if reference and parsed:
                lags.append((reference - parsed).days)
    return {
        "mode": "outlet",
        "rule": "Each outlet's sales are counted up to and including its own last consumable posting date.",
        "salesLastDate": sales_last_date,
        "consumableLastPosting": consumable_last_date,
        "outletsWithPosting": len(cutoffs),
        "outletsWithoutPosting": sorted(activity_codes - set(cutoffs)),
        "outletsTruncated": len(truncated),
        "truncatedOutlets": truncated[:50],
        "maxLagDays": max(lags) if lags else None,
        "medianLagDays": sorted(lags)[len(lags) // 2] if lags else None,
        "salesTotal": round_money(sales_total),
        "salesBase": round_money(base),
        "excludedSales": round_money(sales_total - base),
        "consumableTotal": round_money(consumable_total),
        # Both readings of the headline number, so the effect of the rule is
        # visible on the dashboard rather than inferred from a moved figure.
        "rateOnFullSales": (consumable_total / sales_total) if sales_total else None,
        "rateOnPostedWindow": (consumable_total / base) if base else None,
    }


def validate_payload(payload: dict[str, Any]) -> None:
    outlets = payload.get("outlets", [])
    daily = payload.get("daily", [])
    targets = payload.get("targets", [])
    if not outlets:
        raise ValueError("No outlets were produced after joining the source files.")
    if not daily:
        raise ValueError("No outlet-day records were produced from the source files.")
    codes = [outlet["code"] for outlet in outlets]
    if len(codes) != len(set(codes)):
        raise ValueError("The normalized outlet table contains duplicate outlet codes.")
    unknown_daily_codes = sorted({row["code"] for row in daily} - set(codes))
    if unknown_daily_codes:
        raise ValueError(f"Daily records contain unknown outlet codes: {', '.join(unknown_daily_codes[:20])}")
    if not any(row["sales"] for row in daily):
        raise ValueError("Overall sales are zero across the complete source period.")
    if not any(row["pnpSales"] for row in daily):
        raise ValueError("PNP Sales are zero. Confirm that Sales-Till still uses Division = FRESH PRODUCE.")
    if not any(row["consumable"] for row in daily):
        raise ValueError("No consumable movement value was found.")
    if not any(row["wastage"] for row in daily):
        raise ValueError("No wastage movement value was found.")
    if any("consumableLastPosting" not in outlet for outlet in outlets):
        raise ValueError("Every outlet must carry consumableLastPosting, even when it is null.")

    window = (payload.get("dataQuality") or {}).get("consumableWindow") or {}
    if not window.get("outletsWithPosting"):
        raise ValueError(
            "No outlet has a consumable posting date, so the consumable denominator would be "
            "empty. Check that the CONSUMABLE export still carries a Posting Date column."
        )
    if not window.get("salesBase"):
        raise ValueError(
            "The consumable sales base is zero. Every outlet's consumable postings fall outside "
            "its sales dates, which means the two exports describe different periods."
        )
    windowed_rate = window.get("rateOnPostedWindow")
    if windowed_rate is not None and not 0 <= windowed_rate < 0.25:
        raise ValueError(
            f"Consumable % on the posted window computed as {windowed_rate:.2%}, outside the "
            f"plausible 0-25% band."
        )
    if not targets or any(
        target.get(field) is None
        for target in targets
        for field in ("consumableTarget", "wastageSalesTarget", "wastagePnpTarget")
    ):
        raise ValueError("One or more benchmark target percentages are missing or invalid.")

    quality = payload.get("dataQuality", {})
    coverage = quality.get("targetSalesCoverage")
    if coverage is not None and coverage < 0.90:
        raise ValueError(
            f"Only {coverage:.1%} of sales belongs to outlets present in the zone master. "
            f"Refresh Zone-Distribution.xlsx before publishing."
        )
    reconciliation = quality.get("sales", {}).get("reconciliation", {})
    if reconciliation.get("matches") is False:
        ratio = reconciliation.get("differenceRatio") or 0
        if ratio > 0.001:
            raise ValueError(
                f"Sales detail rows total {reconciliation['detailTotal']:,.2f} but the file's own "
                f"total row says {reconciliation['reportedTotal']:,.2f} "
                f"({ratio:.3%} apart). Re-export Sales-Till.xlsx."
            )


def build_dashboard_data(input_dir: Path) -> dict[str, Any]:
    files = resolve_source_files(input_dir)

    target_map, target_rows = read_targets(files["targets"])
    zone_map, zone_quality = read_zone_distribution(files["zones"], target_map)
    sales_daily, sales_quality = read_sales(files["sales"])
    consumable_daily, consumable_names, consumable_quality, consumable_detail = read_sap_text_export(files["consumable"])
    wastage_daily, wastage_names, wastage_quality, wastage_detail = read_sap_text_export(files["wastage"])
    material_catalog, material_daily = build_material_view(
        consumable_detail, wastage_detail
    )
    # Denominator rule for Consumable % on Sales: each outlet's sales are
    # counted only up to its own last consumable posting date.
    consumable_cutoffs = consumable_posting_cutoffs(consumable_daily)

    activity_codes = {
        code for code, _ in set(sales_daily) | set(consumable_daily) | set(wastage_daily)
    }
    all_codes = sorted(set(zone_map) | activity_codes)
    outlets: list[dict[str, Any]] = []
    for code in all_codes:
        if code in zone_map:
            outlet = dict(zone_map[code])
        else:
            outlet = {
                "code": code,
                "name": first_present([consumable_names.get(code, ""), wastage_names.get(code, ""), code]),
                "regionalLeader": "Unmapped",
                "regionalLeaderFullName": "",
                "zone": "Unmapped",
                "zonalFullName": "",
                "division": "Unmapped",
                "district": "Unmapped",
                "area": "Unmapped",
                "format": "Unmapped",
                "pnpStatus": "Unmapped",
                "ownership": "Unmapped",
                "locationType": "",
                "locationClass": "",
                "populationDensity": "",
                "incomeLevel": "",
                "launchDate": None,
                "sft": None,
                "criteria": "Unmapped",
                "consumableTarget": None,
                "wastageSalesTarget": None,
                "wastagePnpTarget": None,
                "mapped": False,
            }
        outlet["hasActivity"] = code in activity_codes
        # None means this outlet posted no consumable movement in the loaded
        # period, so its consumable rate is unknown rather than zero.
        outlet["consumableLastPosting"] = consumable_cutoffs.get(code)
        outlets.append(outlet)

    all_daily_keys = sorted(set(sales_daily) | set(consumable_daily) | set(wastage_daily), key=lambda item: (item[1], item[0]))
    daily_rows: list[dict[str, Any]] = []
    for code, date in all_daily_keys:
        sales = sales_daily.get((code, date), {})
        daily_rows.append(
            {
                "date": date.isoformat(),
                "code": code,
                "sales": round_money(sales.get("sales", 0.0)),
                "pnpSales": round_money(sales.get("pnpSales", 0.0)),
                "consumable": round_money(consumable_daily.get((code, date), 0.0)),
                "wastage": round_money(wastage_daily.get((code, date), 0.0)),
            }
        )

    consumable_totals: dict[str, float] = defaultdict(float)
    wastage_totals: dict[str, float] = defaultdict(float)
    for (code, _), value in consumable_daily.items():
        consumable_totals[code] += value
    for (code, _), value in wastage_daily.items():
        wastage_totals[code] += value

    source_max_dates = [
        parse_date(sales_quality["dateMax"]),
        parse_date(consumable_quality["dateMax"]),
        parse_date(wastage_quality["dateMax"]),
    ]
    source_max_dates = [value for value in source_max_dates if value]
    sales_total_by_code: dict[str, float] = defaultdict(float)
    for (code, _), row in sales_daily.items():
        sales_total_by_code[code] += row["sales"]
    mapped_sales = sum(value for code, value in sales_total_by_code.items() if code in zone_map)
    total_sales = sum(sales_total_by_code.values())

    data_quality = {
        "activeOutlets": len(activity_codes),
        "mappedActiveOutlets": len(activity_codes & set(zone_map)),
        "unmappedActiveOutlets": sorted(activity_codes - set(zone_map)),
        "mappedWithoutActivity": sorted(set(zone_map) - activity_codes),
        "negativeNetConsumableOutlets": sorted(
            code for code, value in consumable_totals.items() if value < 0
        ),
        "negativeNetWastageOutlets": sorted(
            code for code, value in wastage_totals.items() if value < 0
        ),
        "targetMappedOutlets": sum(
            1 for outlet in outlets if outlet["mapped"] and outlet["consumableTarget"] is not None
        ),
        "targetSalesCoverage": mapped_sales / total_sales if total_sales else None,
        "consumableWindow": build_consumable_window(
            daily_rows,
            consumable_cutoffs,
            activity_codes,
            sales_quality["dateMax"],
            consumable_quality["dateMax"],
        ),
        "periodAlignment": build_period_alignment(
            {
                "sales": sales_quality,
                "consumable": consumable_quality,
                "wastage": wastage_quality,
            }
        ),
        "zone": zone_quality,
        "sales": sales_quality,
        "consumable": consumable_quality,
        "wastage": wastage_quality,
    }

    payload = {
        # 2: adds dataQuality.periodAlignment, sales reconciliation and SAP
        #    column-mapping diagnostics.
        # 3: adds materials and the calendar.
        # 4: adds outlets[].consumableLastPosting and
        #    dataQuality.consumableWindow. The consumable denominator now
        #    depends on that field, so an older payload would produce an empty base.
        # 5: replaces fixed material leaders with complete outlet/date detail,
        #    so material rankings can follow every global filter.
        "schemaVersion": 5,
        "generatedAt": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
        "asOf": iso_date(min(source_max_dates) if source_max_dates else None),
        "dateRange": {
            "min": sales_quality["dateMin"],
            "max": sales_quality["dateMax"],
        },
        "metricDefinitions": {
            "pnpSales": "Sales-Till Division = FRESH PRODUCE",
            "consumableRate": (
                "Net Consumable Value / Overall POS NSI counted up to each outlet's "
                "last consumable posting date"
            ),
            "consumableWindow": (
                "Consumable issues post in batches, so each outlet's sales base stops on the day "
                "of its last consumable posting. An outlet that posted nothing in the period has "
                "no consumable rate rather than 0%."
            ),
            "wastageSalesRate": "Net Wastage Value / Overall POS NSI",
            "wastagePnpRate": "Net Wastage Value / FRESH PRODUCE POS NSI",
            "sapNetValue": "Movement issues less reversals, using Posting Date",
            "aggregateRate": "Sum of value divided by sum of denominator (weighted performance)",
            "outletAverage": "Simple average of valid outlet-level percentages within the selected group",
        },
        "sourceFiles": [
            {"id": "sales", "name": files["sales"].name, "rows": sales_quality["rows"], "dateMin": sales_quality["dateMin"], "dateMax": sales_quality["dateMax"]},
            {"id": "consumable", "name": files["consumable"].name, "rows": consumable_quality["rows"], "dateMin": consumable_quality["dateMin"], "dateMax": consumable_quality["dateMax"]},
            {"id": "wastage", "name": files["wastage"].name, "rows": wastage_quality["rows"], "dateMin": wastage_quality["dateMin"], "dateMax": wastage_quality["dateMax"]},
            {"id": "zones", "name": files["zones"].name, "rows": zone_quality["rows"], "dateMin": None, "dateMax": None},
            {"id": "targets", "name": files["targets"].name, "rows": len(target_rows), "dateMin": None, "dateMax": None},
        ],
        "targets": target_rows,
        "outlets": outlets,
        "daily": daily_rows,
        # [date, isoWeek, isoWeekday] for every date present in the daily rows.
        "calendar": build_calendar(sorted({row["date"] for row in daily_rows})),
        # [code, description, unit]; material rows below reference these by index.
        "materialCatalog": material_catalog,
        # kind -> [[postingDate, outletCode, materialId, netValue, netQuantity], ...]
        "materialDaily": material_daily,
        "dataQuality": data_quality,
    }
    validate_payload(payload)
    return payload


def write_dashboard_data(input_dir: Path, output_path: Path) -> dict[str, Any]:
    payload = build_dashboard_data(input_dir)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(payload, ensure_ascii=False, separators=(",", ":"), allow_nan=False),
        encoding="utf-8",
    )
    return payload


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=Path(".drive-data"))
    parser.add_argument("--output", type=Path, default=Path("dist/data/dashboard-data.json"))
    args = parser.parse_args()
    result = write_dashboard_data(args.input, args.output)
    print(
        json.dumps(
            {
                "status": "ok",
                "outlets": len(result["outlets"]),
                "dailyRows": len(result["daily"]),
                "asOf": result["asOf"],
                "output": str(args.output),
            }
        )
    )
